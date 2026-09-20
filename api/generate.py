"""
Vercel Serverless Function — Gerador APR Suape
"""

import json
import base64
import os
from io import BytesIO
from datetime import datetime
from http.server import BaseHTTPRequestHandler
import openpyxl
from openpyxl.cell.cell import MergedCell

# Path do template (será upload em Vercel)
TEMPLATE_PATH = '/tmp/APR_TEMPLATE.xlsx'

def formatar_data(data_iso):
    if not data_iso:
        return ''
    try:
        d = datetime.fromisoformat(data_iso)
        return d.strftime('%d/%m/%Y')
    except:
        return data_iso

def preencher_apr_suape(dados, template_bytes):
    """Preenche template preservando formatação"""
    try:
        # Carregar template de bytes
        wb = openpyxl.load_workbook(BytesIO(template_bytes))
        ws = wb['Plan1']
        
        data_inicio = formatar_data(dados.get('dataInicio', ''))
        data_final = formatar_data(dados.get('dataFinal', ''))
        
        preenchimentos = [
            ('A4', f"Contrato Nº: {dados.get('contrato', '0000')}"),
            ('D4', f"Empresa: {dados.get('empresa', 'Decal Brasil Ltda')}"),
            ('G4', f"Colocar a unidade Operacional: {dados.get('unidade', 'PC-03')}"),
            ('A6', f"Processo: {dados.get('processo', '')}"),
            ('D6', f"Função: {dados.get('funcao', '')}"),
            ('G6', f"Data de Início da atividade: {data_inicio}"),
            ('I6', f"Data Final da atividade: {data_final}"),
            ('K6', f"Número da APR/Revisão: {dados.get('aprNumero', '')}"),
            ('A11', f"Descrição da Atividade: {dados.get('descricao', '')}"),
        ]
        
        for cell_ref, valor in preenchimentos:
            cell = ws[cell_ref]
            if isinstance(cell, MergedCell):
                for merged_range in ws.merged_cells.ranges:
                    if cell_ref in merged_range:
                        master_cell = ws[merged_range.min_row, merged_range.min_col]
                        master_cell.value = valor
                        break
            else:
                cell.value = valor
        
        if dados.get('responsavel'):
            ws['A10'] = dados.get('responsavel', '')
        
        if dados.get('membros'):
            membros_str = dados.get('membros', '')
            if isinstance(membros_str, list):
                membros_str = ', '.join(membros_str)
            ws['G10'] = membros_str
        
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
    
    except Exception as e:
        raise Exception(f"Erro ao preencher: {str(e)}")

class handler(BaseHTTPRequestHandler):
    
    def do_OPTIONS(self):
        self.send_response(204)
        self.send_header('Access-Control-Allow-Origin', '*')
        self.send_header('Access-Control-Allow-Methods', 'GET,OPTIONS,PATCH,DELETE,POST,PUT')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type')
        self.end_headers()
    
    def do_POST(self):
        try:
            # Ler body
            content_length = int(self.headers.get('Content-Length', 0))
            body = self.rfile.read(content_length)
            dados = json.loads(body.decode('utf-8'))
            
            # Carregar template (será feito upload em Vercel ou servido)
            # Por enquanto, usar template embutido em base64
            template_b64 = os.getenv('TEMPLATE_BASE64', '')
            if not template_b64:
                raise Exception('Template não configurado')
            
            template_bytes = base64.b64decode(template_b64)
            
            # Preencher
            excel_bytes = preencher_apr_suape(dados, template_bytes)
            
            # Retornar base64
            b64 = base64.b64encode(excel_bytes).decode('utf-8')
            
            response = {
                'success': True,
                'base64': b64,
                'filename': f"APR_{dados.get('contrato', 'GERADA')}_{datetime.now().strftime('%Y%m%d')}.xlsx"
            }
            
            # Headers
            self.send_response(200)
            self.send_header('Content-Type', 'application/json')
            self.send_header('Access-Control-Allow-Origin', '*')
            self.end_headers()
            
            self.wfile.write(json.dumps(response).encode('utf-8'))
        
        except Exception as e:
            response = {'success': False, 'error': str(e)}
            self.send_response(500)
            self.send_header('Content-Type', 'application/json')
            self.send_header('Access-Control-Allow-Origin', '*')
            self.end_headers()
            self.wfile.write(json.dumps(response).encode('utf-8'))
