"""
Gerador APR Suape — Orbit 360
Refatorado para ler template Excel local (sem base64)
"""

import os
import io
from flask import Flask, jsonify, request
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter

app = Flask(__name__)

# Path do template (relativo ao root do projeto)
TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), '../template_apr.xlsx')

@app.route('/', methods=['GET'])
def home():
    return jsonify({
        "status": "ok",
        "message": "Gerador APR Suape — Orbit 360",
        "api": "/api/generate",
        "version": "1.1.0"
    })

@app.route('/api/generate', methods=['POST'])
def generate_apr():
    """
    Gera APR Excel preenchido
    Body JSON:
    {
        "contrato": "valor",
        "empresa": "valor",
        ...
    }
    """
    try:
        data = request.json or {}
        
        # Valida dados mínimos
        if not data.get('empresa'):
            return jsonify({"error": "Campo 'empresa' obrigatório"}), 400
        
        # Carrega template
        if not os.path.exists(TEMPLATE_PATH):
            return jsonify({"error": f"Template não encontrado em {TEMPLATE_PATH}"}), 500
        
        wb = load_workbook(TEMPLATE_PATH)
        ws = wb.active
        
        # Mapeamento de campos → células (conforme template oficial)
        field_mapping = {
            'contrato': 'A4',
            'empresa': 'D4',
            'unidade': 'G4',
            'processo': 'A6',
            'funcao': 'D6',
            'data_inicio': 'G6',
            'data_final': 'I6',
            'apr_numero': 'K6',
            'engenheiro': 'A10',
            'membros': 'G10',
            'descricao': 'A11',
        }
        
        # Preenche células
        for field, cell in field_mapping.items():
            if field in data and data[field]:
                ws[cell].value = data[field]
        
        # Gera arquivo em memória
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        # Retorna como base64 (simples)
        import base64
        file_b64 = base64.b64encode(output.getvalue()).decode('utf-8')
        
        return jsonify({
            "status": "success",
            "file_b64": file_b64,
            "filename": f"APR_{data.get('empresa', 'suape')}.xlsx"
        }), 200
    
    except Exception as e:
        return jsonify({"error": str(e)}), 500

if __name__ == '__main__':
    app.run(debug=False, host='0.0.0.0', port=5000)
